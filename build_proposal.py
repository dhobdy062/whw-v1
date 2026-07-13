from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = "/Users/donhobdyjr/Documents/Warren Realty Project/warren-hall-digital-presence-proposal.docx"

NAVY = RGBColor(18, 39, 61)
INK = RGBColor(39, 47, 56)
MUTED = RGBColor(103, 113, 124)
GOLD = RGBColor(176, 137, 70)
PALE = "F5F2EC"
LIGHT_BLUE = "EAF0F4"
LIGHT_GOLD = "F4EBDD"
WHITE = RGBColor(255, 255, 255)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ['val', 'sz', 'space', 'color']:
                if key in kwargs[edge]:
                    element.set(qn('w:{}'.format(key)), str(kwargs[edge][key]))

def set_run_font(run, name='Aptos', size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def set_para(p, before=0, after=6, line=1.12, align=None, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if keep:
        pf.keep_with_next = True

def add_text(doc, text='', size=11, color=INK, bold=False, italic=False, align=None, before=0, after=6, line=1.12, style=None):
    p = doc.add_paragraph(style=style)
    set_para(p, before, after, line, align)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_para(p, before=18 if level == 1 else 12, after=7, line=1.0, keep=True)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 12.5, color=NAVY if level == 1 else GOLD, bold=True)
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.28 + 0.22 * level)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    set_para(p, after=4, line=1.10)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.left_indent = Inches(0.32)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    set_para(p, after=5, line=1.10)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    return p

def add_callout(doc, label, text, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.25)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=160, start=180, bottom=160, end=180)
    set_cell_border(cell, top={'val':'single','sz':'6','color':'D9D1C4'}, bottom={'val':'single','sz':'6','color':'D9D1C4'}, left={'val':'single','sz':'6','color':'D9D1C4'}, right={'val':'single','sz':'6','color':'D9D1C4'})
    p = cell.paragraphs[0]
    set_para(p, after=3, line=1.08)
    r = p.add_run(label.upper() + "  ")
    set_run_font(r, size=9, color=GOLD, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=11, color=NAVY, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.width = Inches(widths[i])
        set_cell_shading(c, '18344D')
        set_cell_margins(c, top=110, start=130, bottom=110, end=130)
        p = c.paragraphs[0]
        set_para(p, after=0, line=1.0)
        r = p.add_run(h)
        set_run_font(r, size=9.5, color=WHITE, bold=True)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            c.width = Inches(widths[i])
            set_cell_shading(c, 'FFFFFF' if ridx % 2 == 0 else 'F7F8F9')
            set_cell_margins(c, top=105, start=130, bottom=105, end=130)
            set_cell_border(c, bottom={'val':'single','sz':'4','color':'D7DDE2'})
            p = c.paragraphs[0]
            set_para(p, after=0, line=1.08)
            r = p.add_run(val)
            set_run_font(r, size=9.5, color=INK, bold=(i == 0 and len(row) > 2))
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table

def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

def set_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    set_para(p, after=0, line=1.0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run('WARREN HALL  |  DIGITAL PRESENCE & GROWTH PROPOSAL')
    set_run_font(r, size=8, color=MUTED, bold=True)
    footer = section.footer
    p = footer.paragraphs[0]
    set_para(p, before=0, after=0, line=1.0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('Prepared for Warren Hall  •  July 2026')
    set_run_font(r, size=8, color=MUTED)

def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(0.92)
    sec.right_margin = Inches(0.92)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)
    set_header_footer(sec)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    for sname in ['List Bullet', 'List Number']:
        st = styles[sname]
        st.font.name = 'Aptos'
        st._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
        st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
        st.font.size = Pt(10.5)

    # Cover / proposal centerpiece
    add_text(doc, 'DIGITAL PRESENCE & GROWTH PROPOSAL', size=10, color=GOLD, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=34, after=10, line=1.0)
    add_text(doc, 'A stronger digital platform\nfor Warren Hall', size=27, color=NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=10, line=1.0)
    add_text(doc, 'Website launch • Luxury / move-up positioning • Lead generation • Social content engine', size=12, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=24, line=1.0)
    add_callout(doc, 'The opportunity', 'Turn your existing production credibility into a modern, measurable system that attracts more $750K+ buyers and sellers—and makes it easier for every lead to take the next step.', fill=LIGHT_GOLD)
    add_text(doc, 'Prepared for', size=9, color=MUTED, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=20, after=2)
    add_text(doc, 'Warren Hall', size=17, color=NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_text(doc, 'Compass • Raleigh / Triangle market', size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=22)
    add_text(doc, 'Prepared by Retrospxt Holdings - Don Hobdy Jr', size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    add_text(doc, 'July 2026', size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    add_page_break(doc)

    add_heading(doc, '1. Executive summary', 1)
    add_text(doc, 'You already have the production history, local credibility, and client experience to compete for higher-value business. The gap is not capability—it is the digital infrastructure that presents that capability consistently, captures demand, and follows up when you are busy serving clients.', after=8)
    add_text(doc, 'This proposal starts with the website because it becomes the owned destination for every social post, referral, listing video, market update, and paid campaign. From there, we build a repeatable content, training, and lead-management system around your voice and expertise.', after=8)
    add_callout(doc, 'Recommended direction', 'Raleigh native. Triangle real estate advisor. Clear strategy for buying, selling, and moving up.', fill=LIGHT_BLUE)
    add_heading(doc, 'What this creates for you', 2)
    for t in [
        'A credible, personal-brand website designed for buyers and sellers—not a generic brokerage page.',
        'A clear luxury and move-up lane beginning around $750,000, supported by your broader Triangle experience.',
        'A content engine that turns one approved idea into Instagram, TikTok/Reels, Pinterest, website, and email assets.',
        'Lead capture and follow-up paths that make it easier to convert attention into consultations.',
        'A measurable foundation for future advertising, referrals, IDX/search, and CRM automation.'
    ]: add_bullet(doc, t)
    add_heading(doc, '2. Why now', 1)
    add_text(doc, 'Your direct discovery conversation identified three timing factors:', after=5)
    for t in [
        'You recently changed companies and do not yet have a new website.',
        'You want to expand your direct business into the $700K–$750K+ segment.',
        'You need new lead sources and want social video and automation to reduce the friction of creating content personally.'
    ]: add_bullet(doc, t)
    add_text(doc, 'The market also rewards better positioning. Raleigh has more inventory and more price sensitivity than the peak-pandemic environment, so sellers need strategy and buyers need clarity. A polished platform helps you communicate both.', after=8)
    add_page_break(doc)

    add_heading(doc, '3. Phase 1 scope: website-first launch', 1)
    add_text(doc, 'Phase 1 is a 6-week launch designed to establish the core digital asset before scaling social publishing and paid lead generation.', after=8)
    add_heading(doc, 'Website experience', 2)
    for t in [
        'Personal-brand homepage with strong photography, proof, service area, and two primary calls to action.',
        'Buy, Sell, and Luxury & Move-Up pathways.',
        'About page with your story, Raleigh roots, production proof, service philosophy, and approved credentials.',
        'Results / Portfolio section with 6–10 approved representative transactions or case studies.',
        'Neighborhood hub with initial Raleigh/Triangle pages and a structure that can expand over time.',
        'Compass or approved IDX/home-search handoff, depending on brokerage and MLS requirements.',
        'Confidential home valuation form and consultation scheduling.',
        'Insights/blog structure for market updates, neighborhood content, and repurposed social content.',
        'Privacy, consent, brokerage disclosure, compliance, analytics, and mobile QA.'
    ]: add_bullet(doc, t)
    add_heading(doc, 'Lead-generation foundation', 2)
    add_table(doc, ['Audience', 'Offer / CTA', 'Business value'], [
        ['Move-up / luxury seller', 'Request a private valuation', 'Creates a high-intent seller conversation'],
        ['Luxury buyer', 'Curated $750K+ home search', 'Captures preferences and timing'],
        ['Relocation lead', 'Triangle relocation brief', 'Turns local expertise into an email capture'],
        ['Referral partner', 'Refer a client to you', 'Creates a simple referral path'],
    ], [1.55, 2.15, 2.55])
    add_text(doc, 'All listing claims, brokerage language, volume claims, awards, and transaction proof will be confirmed with you and Compass before publication.', size=9.5, color=MUTED, italic=True, before=6, after=4)
    add_heading(doc, 'Lead nurturing & follow-up system', 2)
    add_text(doc, 'The website should not simply collect contact information. Each inquiry should enter a clear follow-up path based on what the person needs, how soon they are moving, and how engaged they are.', after=6)
    add_table(doc, ['Lead segment', 'Initial response', 'Nurture path'], [
        ['Active buyer or seller', 'Immediate confirmation + booking link', 'Personal follow-up task, consultation scheduling, next-step checklist'],
        ['Future buyer or seller', 'Helpful guide + market context', 'Education sequence, periodic market updates, re-engagement prompts'],
        ['Relocation or referral', 'Personal welcome + local resource', 'Area guide, vendor/referral support, consultation invitation'],
    ], [1.65, 2.2, 2.4])
    add_text(doc, 'Recommended baseline: respond immediately by email, create a same-day personal follow-up task, and continue useful—not aggressive—nurture until the lead books, opts out, or is marked inactive. The exact email/SMS cadence and compliance language will be finalized with your CRM and Compass requirements.', size=9.5, color=MUTED, italic=True, before=6, after=4)
    add_page_break(doc)

    add_heading(doc, '4. Content and automation system', 1)
    add_text(doc, 'After the website is live, the social system will make your expertise visible without requiring you to manually create every post from scratch.', after=8)
    add_heading(doc, 'Primary platforms', 2)
    add_table(doc, ['Platform', 'Role', 'Starting cadence'], [
        ['Instagram', 'Primary trust, local authority, listings, Reels, Stories', '3 feed posts / Reels weekly + Stories'],
        ['TikTok / Reels', 'Reach, education, personality, market hooks', '3 short videos weekly'],
        ['Pinterest', 'Evergreen discovery for neighborhoods, design, buyers and sellers', '5–10 Pins weekly'],
    ], [1.35, 3.55, 1.35])
    add_heading(doc, 'Content pillars', 2)
    for t in [
        'Raleigh luxury market intelligence: pricing, inventory, negotiation, and timing.',
        'Property storytelling: architecture, design, floor plans, renovations, and lifestyle fit.',
        'Neighborhood intelligence: Raleigh and Triangle communities, amenities, commute context, and local perspective.',
        'Buyer and seller education: inspections, appraisal, preparation, offers, relocation, and common mistakes.',
        'Your trust story: client experience, behind the scenes, vendor network, and local relationships.'
    ]: add_bullet(doc, t)
    add_heading(doc, 'One-to-many production model', 2)
    add_text(doc, 'One approved weekly source idea—usually a 60–90 second video, market brief, or property story—can become a Reel, TikTok, carousel, Story sequence, Pinterest Pin, website insight, and email snippet. Automation supports drafting, formatting, scheduling, UTM tracking, and reporting. You retain final approval over claims, listings, prices, client stories, and synthetic/AI-assisted representations.', after=8)
    add_callout(doc, 'Pilot recommendation', 'Start with a defined AI/video pilot: a controlled source session, a small batch of approved virtual-tour and market-video formats, and clear review/disclosure rules.', fill=LIGHT_GOLD)
    add_heading(doc, 'Longer-term platform opportunity', 2)
    add_text(doc, 'You have expressed interest in shadowing the type of integrated system Joe is building for Carolina Insurance Advisors, including AI videos for social media, training, and future lead-generation concepts. That opportunity should be explored after the Phase 1 foundation is stable so the work remains practical, measurable, and reusable.', after=7)
    for t in [
        'AI-assisted property tours and market explainers for social media.',
        'Training videos for your team, referral partners, or future agents.',
        'A lead-generation concept based on your Raleigh knowledge and your broader idea for a tool valuable in any market.',
        'A shared dashboard for content performance, lead engagement, response tasks, and pipeline visibility.'
    ]: add_bullet(doc, t)
    add_callout(doc, 'Expansion principle', 'Phase 1 proves the audience, message, website, and lead path. Phase 2 turns the best-performing pieces into a scalable real-estate growth product.', fill=LIGHT_BLUE)

    add_heading(doc, '5. Timeline', 1)
    add_text(doc, 'Target launch window: approximately 6 weeks from receipt of required materials and approvals.', after=8)
    add_table(doc, ['Timing', 'Focus', 'Key outputs'], [
        ['Week 1', 'Discovery & positioning', 'Goals, audiences, proof inventory, brokerage/Compass requirements, sitemap, CTA decisions'],
        ['Week 2', 'Brand & content architecture', 'Visual direction, messaging, page outlines, content pillars, lead magnet outline'],
        ['Week 3', 'Design', 'Homepage and key-page designs, mobile direction, CTA and form experience'],
        ['Week 4', 'Build & copy', 'Responsive site build, copy, portfolio, neighborhood structure, analytics setup'],
        ['Week 5', 'Lead & content systems', 'Forms, scheduling, CRM handoff, initial content templates, social launch calendar'],
        ['Week 6', 'QA & launch', 'Mobile testing, form tests, compliance review, SEO basics, launch, handoff, 30-day backlog'],
    ], [1.0, 1.65, 3.6])
    add_heading(doc, 'Client inputs needed to stay on schedule', 2)
    for t in [
        'Approved headshots, property photography, video, logo/brand assets, and contact details.',
        'Compass brokerage wording, BIC requirements, license information, and approved claims.',
        'Six to ten representative transactions, testimonials, awards, and volume documentation.',
        'Preferred service area, neighborhoods, price range, CRM, domain, and current lead sources.',
        'One focused approval meeting per week during the build.'
    ]: add_bullet(doc, t)
    add_page_break(doc)

    add_heading(doc, '6. Investment', 1)
    add_text(doc, 'The initial setup is intentionally discounted to establish the relationship, launch the platform, and generate enough real performance data to price the ongoing marketing work responsibly. Software, hosting, email, CRM, scheduling, media, IDX, paid advertising, and high-touch video production may be separate depending on the selected stack and approved scope.', after=8)
    add_table(doc, ['Commercial component', 'Investment', 'Purpose'], [
        ['One-time setup', '$1,200', 'Discounted onboarding, initial platform setup, lead capture, analytics, core configuration, and launch support'],
        ['Months 1–3 platform fee', '$750 / month', 'Website/platform maintenance, reporting, lead-routing support, nurture operations, and agreed social/marketing support'],
        ['After 90-day review', '$750–$1,500 / month target range', 'Adjusted fixed marketing fee based on measured engagement, lead quality, attributable opportunity, and workload'],
    ], [1.65, 1.55, 3.0])
    add_callout(doc, 'Recommended starting engagement', 'Begin with a $1,200 setup fee plus a $750 monthly marketing/platform fee for the first three months.', fill=LIGHT_BLUE)
    add_heading(doc, '90-day marketing agreement review', 2)
    add_text(doc, 'At the end of month three, we will review the marketing agreement using a documented scorecard. The goal is to set a fair fixed monthly marketing fee for the next period—not to charge a percentage of commission income.', after=6)
    add_table(doc, ['Review input', 'What will be measured'], [
        ['Website performance', 'Traffic sources, organic search, conversion rate, form submissions, consultations, valuation requests'],
        ['Social performance', 'Reach, saves, shares, profile visits, link clicks, content output, and engagement quality'],
        ['Lead quality', 'Qualified conversations, response time, booked appointments, active opportunities, and source attribution'],
        ['Business impact', 'Publicly available market data, attributable sales or opportunities, client value, and revenue context'],
        ['Service workload', 'Content volume, platform maintenance, reporting, nurture operations, and support requirements'],
    ], [1.65, 4.55])
    add_text(doc, 'Sales attributed to the website, organic search, or social media will be considered as evidence of platform value and return on investment, but will not be used as a commission split or percentage-of-income formula. Any future change to the fixed marketing fee will be agreed in writing before the next billing period.', size=9.5, color=MUTED, italic=True, before=6, after=4)
    add_heading(doc, '7. Business benefits and success measures', 1)
    add_text(doc, 'The project is designed to create business leverage, not just attractive marketing.', after=8)
    add_table(doc, ['Business challenge', 'What the project changes', 'Expected benefit'], [
        ['No current website after company change', 'Launch a credible owned destination', 'More trust, control, and conversion from every channel'],
        ['Desire to expand into $750K+ business', 'Build a visible luxury / move-up lane', 'Better alignment between message and target opportunity'],
        ['Limited time for social production', 'Create one-to-many content workflows', 'More consistent visibility with less personal friction'],
        ['Need new lead sources', 'Add offers, forms, tracking, and follow-up', 'More measurable conversations outside purchased leads'],
        ['Leads can be difficult to maintain', 'Connect capture to CRM and response tasks', 'Fewer missed opportunities and clearer pipeline visibility'],
    ], [1.65, 2.55, 2.05])
    add_heading(doc, '90-day success measures', 2)
    for t in [
        'Your website is live, mobile-friendly, indexed, measurable, and passes test leads through the chosen response path.',
        'At least three repeatable content formats are publishing across Instagram, TikTok/Reels, and Pinterest.',
        'Baseline is established for qualified inquiries, consultation bookings, valuation requests, and lead response time.',
        'You have a 30-day content backlog and a simple approval process you can sustain.',
        'The team can identify which content and channels produce qualified conversations—not just views or followers.'
    ]: add_bullet(doc, t)
    add_callout(doc, 'The real outcome', 'A digital system that compounds your reputation: every new listing, client win, market insight, and referral has a place to live, a way to be distributed, and a path back to a conversation.', fill=LIGHT_GOLD)
    add_heading(doc, '8. Next steps', 1)
    for t in [
        'Confirm current Compass affiliation, preferred business name, service area, and approved proof points.',
        'Schedule a 60-minute discovery and asset-collection session.',
        'Approve the visual direction inspired by the List It With Liza reference: personal, editorial, local, and easy to navigate.',
        'Finalize Phase 1 scope, investment, and launch date.',
        'Begin Week 1 discovery.'
    ]: add_number(doc, t)
    add_text(doc, 'Thank you for the opportunity to build this with you.', size=12, color=NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=22, after=3)
    add_text(doc, 'Retrospxt Holdings - Don Hobdy Jr', size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

    doc.save(OUT)

if __name__ == '__main__':
    build()
